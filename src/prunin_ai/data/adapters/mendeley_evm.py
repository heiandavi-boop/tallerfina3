from __future__ import annotations
import csv
from itertools import chain
from pathlib import Path
import numpy as np
import pandas as pd
from ..common import normalize_columns, first_existing, numeric_series, text_series, scale_to_unit_interval


def _read_table(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() == ".csv":
        return pd.read_csv(p)
    if p.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(p, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        source_rows = sheet.iter_rows(values_only=True)
        first_row = next(source_rows, None)
        if first_row and len(first_row) == 1 and isinstance(first_row[0], str) and "," in first_row[0]:
            rows = [next(csv.reader([first_row[0]]))]
            rows.extend(next(csv.reader([str(row[0] or "")])) for row in source_rows)
            workbook.close()
            if not rows:
                raise ValueError(f"El archivo no contiene filas: {p}")
            return pd.DataFrame(rows[1:], columns=rows[0])
        workbook.close()
    return pd.read_excel(p)


def _project_id(df: pd.DataFrame) -> pd.Series:
    col = first_existing(df, ["project_id", "projectid", "project", "id"])
    if not col:
        raise ValueError("No se pudo identificar project_id en el archivo.")
    return df[col].astype(str).str.strip()


def _duration_weeks(df: pd.DataFrame) -> pd.Series:
    col = first_existing(df, [
        "planned_duration_weeks", "planned_duration_week", "planned_duration_days",
        "planned_duration_months", "estimated_timeline_months", "planned_duration",
        "duration"
    ])
    if not col:
        return pd.Series(np.nan, index=df.index)
    x = pd.to_numeric(df[col], errors="coerce")
    if "day" in col:
        return x / 7.0
    if "month" in col:
        return x * 4.345
    return x


def _duration_to_days(values: pd.Series, column_name: str) -> pd.Series:
    x = pd.to_numeric(values, errors="coerce")
    name = column_name.lower()
    if "week" in name:
        return x * 7.0
    if "month" in name:
        return x * 30.4375
    return x


def _delay_days(df: pd.DataFrame, static: pd.DataFrame | None = None) -> pd.Series:
    col = first_existing(df, [
        "schedule_delay_days", "delay_days", "schedule_delay_weeks", "delay_weeks",
        "schedule_delay_months", "delay_months"
    ])
    if col:
        return _duration_to_days(df[col], col).clip(lower=0)

    actual_col = first_existing(df, ["actual_duration_days", "actual_duration_weeks", "actual_duration_months", "actual_duration"])
    planned_col = first_existing(df, [
        "planned_duration_days", "planned_duration_weeks", "planned_duration_months",
        "estimated_timeline_months", "planned_duration"
    ])
    if actual_col and planned_col:
        actual = _duration_to_days(df[actual_col], actual_col)
        planned = _duration_to_days(df[planned_col], planned_col)
        return (actual - planned).clip(lower=0)

    generic_col = first_existing(df, ["schedule_delay", "delay"])
    if generic_col:
        return pd.to_numeric(df[generic_col], errors="coerce").clip(lower=0)
    return pd.Series(np.nan, index=df.index)


def prepare_mendeley(static_path: str | Path, timeseries_path: str | Path, summary_path: str | Path) -> pd.DataFrame:
    static = normalize_columns(_read_table(static_path))
    ts = normalize_columns(_read_table(timeseries_path))
    summary = normalize_columns(_read_table(summary_path))

    static["project_id"] = _project_id(static)
    ts["project_id"] = _project_id(ts)
    summary["project_id"] = _project_id(summary)

    s = pd.DataFrame(index=static.index)
    s["project_id"] = static["project_id"]
    s["planned_duration_weeks"] = _duration_weeks(static)
    s["planned_duration_days"] = s["planned_duration_weeks"] * 7.0
    s["planned_budget"] = numeric_series(static, [
        "budget", "project_budget_usd", "planned_budget", "bac", "baseline_budget"
    ])
    s["baseline_scope_units"] = numeric_series(static, ["baseline_scope_units", "scope_units", "story_points"])
    s["sector"] = text_series(static, ["sector", "industry", "domain"])
    s["project_type"] = text_series(static, ["project_type", "type"])
    s["methodology"] = text_series(static, [
        "methodology", "methodology_used", "delivery_methodology", "execution_methodology"
    ])
    s["complexity"] = text_series(static, ["complexity", "project_complexity"])
    s["criticality"] = text_series(static, ["criticality", "project_criticality"])

    # Sólo se usan equivalencias semánticamente defendibles. Valores no reconocibles quedan NaN.
    s["normalized_risk_exposure_static"] = scale_to_unit_interval(
        numeric_series(static, ["risk_exposure", "risk_score", "risk_history_score", "risk_history"])
    )
    s["governance_health_static"] = scale_to_unit_interval(
        numeric_series(static, ["governance_maturity", "governance_score", "governance_health"])
    )

    t = pd.DataFrame(index=ts.index)
    t["project_id"] = ts["project_id"]
    month_col = first_existing(ts, ["month", "project_month", "period", "time_period", "snapshot_index"])
    t["snapshot_index"] = pd.to_numeric(ts[month_col], errors="coerce") if month_col else ts.groupby("project_id").cumcount() + 1
    t["pv"] = numeric_series(ts, ["pv", "planned_value", "cumulative_pv"])
    t["ev"] = numeric_series(ts, ["ev", "earned_value", "cumulative_ev"])
    t["ac"] = numeric_series(ts, ["ac", "actual_cost", "cumulative_ac"])
    t["spi"] = numeric_series(ts, ["cumulative_spi", "spi_cumulative", "spi"])
    t["cpi"] = numeric_series(ts, ["cumulative_cpi", "cpi_cumulative", "cpi"])

    # Join budget before deriving progress.
    t = t.merge(s[["project_id", "planned_budget"]], on="project_id", how="left")
    t["true_progress"] = t["ev"] / t["planned_budget"].replace(0, np.nan)
    t["true_progress"] = t["true_progress"].clip(lower=0, upper=1.5)
    t["spi"] = t["spi"].where(t["spi"].notna(), t["ev"] / t["pv"].replace(0, np.nan))
    t["cpi"] = t["cpi"].where(t["cpi"].notna(), t["ev"] / t["ac"].replace(0, np.nan))

    # Features V9 sin equivalencia directa permanecen ausentes (NaN).
    for c in [
        "team_utilization", "team_capacity_ratio", "average_productivity", "defect_rate",
        "rework_ratio", "scope_growth_ratio", "dependency_delay_days", "team_health_index"
    ]:
        t[c] = np.nan

    t = t.merge(
        s.drop(columns=["planned_budget"]), on="project_id", how="left"
    )
    t["normalized_risk_exposure"] = t["normalized_risk_exposure_static"]
    t["governance_health_score"] = t["governance_health_static"]
    t = t.drop(columns=["normalized_risk_exposure_static", "governance_health_static"], errors="ignore")

    o = pd.DataFrame(index=summary.index)
    o["project_id"] = summary["project_id"]
    o["delay_days"] = _delay_days(summary)
    cost_overrun_col = first_existing(summary, ["cost_overrun_ratio", "cost_overrun_pct", "cost_overrun_percent"])
    if cost_overrun_col:
        cor = pd.to_numeric(summary[cost_overrun_col], errors="coerce")
        if "pct" in cost_overrun_col or "percent" in cost_overrun_col or cor.dropna().abs().median() > 1:
            cor = cor / 100.0
        o["cost_overrun_ratio"] = cor
    else:
        actual_cost = numeric_series(summary, ["total_cost", "actual_cost", "final_cost"])
        temp = summary.copy()
        temp["project_id"] = summary["project_id"]
        budget_map = s.set_index("project_id")["planned_budget"]
        budget = temp["project_id"].map(budget_map)
        has_actual_cost = actual_cost.notna()
        raw_overrun = numeric_series(summary, ["cost_overrun"])
        amount = raw_overrun.where(raw_overrun.notna(), actual_cost - budget)
        o["cost_overrun_ratio"] = amount / budget.replace(0, np.nan)
        if not has_actual_cost.any() and raw_overrun.isna().all():
            o["cost_overrun_ratio"] = np.nan

    out = t.merge(o, on="project_id", how="left")
    out["data_source"] = "mendeley_2p5sz57wh2_v2"
    out["source_project_id"] = out["project_id"].astype(str)
    out["project_id"] = "mendeley::" + out["source_project_id"]
    out["is_synthetic"] = True
    return out

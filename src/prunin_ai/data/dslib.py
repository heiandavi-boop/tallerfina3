from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
RAW_DIR = ROOT / "data" / "raw" / "dslib"
ARCHIVE = RAW_DIR / "DSLIB3.4.zip"
DATASET_DIR = RAW_DIR / "extracted" / "DSLIB 3.4"
PROJECT_CODE = re.compile(r"^C\d{4}-\d{2}$", re.IGNORECASE)
FILE_PROJECT_CODE = re.compile(r"^(C\d{4}-\d{2})", re.IGNORECASE)
PROJECT_FEATURES = ["planned_duration_days", "planned_budget_bac"]
SNAPSHOT_FEATURES = [
    *PROJECT_FEATURES,
    "elapsed_ratio",
    "percent_complete",
    "ev",
    "pv",
    "ac",
    "spi",
    "cpi",
    "sv",
    "cv",
]


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        return None
    return float(value)


def _datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return None


def _extract_archive() -> None:
    if not ARCHIVE.exists():
        raise FileNotFoundError(f"No existe el archivo DSLIB: {ARCHIVE.relative_to(ROOT)}")
    if DATASET_DIR.exists():
        return
    extraction_root = DATASET_DIR.parent
    extraction_root.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(ARCHIVE) as archive:
        for member in archive.infolist():
            destination = (extraction_root / member.filename).resolve()
            if not destination.is_relative_to(extraction_root.resolve()):
                raise ValueError(f"Ruta inválida dentro del ZIP: {member.filename}")
        archive.extractall(extraction_root)


def _read_project_outcomes() -> dict[str, dict[str, Any]]:
    workbook = load_workbook(DATASET_DIR / "DSLIB_Analysis_Sheet.xlsx", read_only=True, data_only=True)
    sheet = workbook["DSLIB"]
    projects: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row or not isinstance(row[0], str) or not PROJECT_CODE.fullmatch(row[0].strip()):
            continue
        project_id = row[0].strip().upper()
        projects[project_id] = {
            "planned_duration_days": _number(row[11] if len(row) > 11 else None),
            "planned_budget_bac": _number(row[12] if len(row) > 12 else None),
            "actual_duration_days": _number(row[82] if len(row) > 82 else None),
            "actual_cost": _number(row[83] if len(row) > 83 else None),
        }
    workbook.close()
    return projects


def _baseline_dates(path: Path) -> tuple[datetime | None, datetime | None]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Baseline Schedule" not in workbook.sheetnames:
        workbook.close()
        return None, None
    sheet = workbook["Baseline Schedule"]
    indexes: dict[str, int] = {}
    header_row = 0
    for row_number in range(1, min(sheet.max_row, 12) + 1):
        values = [sheet.cell(row=row_number, column=column).value for column in range(1, sheet.max_column + 1)]
        candidate = {str(value).strip().casefold(): index for index, value in enumerate(values) if value is not None}
        if {"id", "baseline start", "baseline end"} <= candidate.keys():
            indexes = candidate
            header_row = row_number
            break
    start = finish = None
    if header_row:
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if row and row[indexes["id"]] == 0:
                start = _datetime(row[indexes["baseline start"]])
                finish = _datetime(row[indexes["baseline end"]])
                break
    workbook.close()
    return start, finish


def _has_baseline_schedule(path: Path) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=True)
    has_baseline = "Baseline Schedule" in workbook.sheetnames
    workbook.close()
    return has_baseline


def _safe_snapshots(path: Path, project_id: str, outcome: dict[str, Any],
                    planned_start: datetime | None, planned_finish: datetime | None) -> list[dict[str, Any]]:
    if planned_start is None or planned_finish is None:
        return []
    planned_duration = outcome["planned_duration_days"]
    bac = outcome["planned_budget_bac"]
    actual_duration = outcome["actual_duration_days"]
    actual_cost = outcome["actual_cost"]
    if planned_duration is None or planned_duration <= 0 or bac is None or bac <= 0:
        return []
    delay = actual_duration - planned_duration if actual_duration is not None and actual_duration >= 0 else None
    overrun = (actual_cost - bac) / bac if actual_cost is not None and actual_cost >= 0 else None

    workbook = load_workbook(path, read_only=True, data_only=True)
    snapshots = []
    for sheet_name in workbook.sheetnames:
        if not sheet_name.casefold().startswith("project control"):
            continue
        sheet = workbook[sheet_name]
        status_date = _datetime(sheet.cell(row=1, column=3).value)
        if status_date is None or status_date >= planned_finish:
            continue
        headers = [sheet.cell(row=4, column=column).value for column in range(1, sheet.max_column + 1)]
        columns = {str(value).strip().casefold(): index for index, value in enumerate(headers) if value is not None}
        required = {
            "ev": columns.get("earned value (ev)"),
            "pv": columns.get("planned value (pv)"),
            "ac": columns.get("actual cost"),
            "progress": columns.get("percentage completed"),
        }
        if any(index is None for index in required.values()):
            continue
        summary = next((row for row in sheet.iter_rows(min_row=5, values_only=True) if row and row[0] == 0), None)
        if summary is None:
            continue
        ev, pv, ac, progress = (_number(summary[required[key]]) for key in ("ev", "pv", "ac", "progress"))
        if any(value is None for value in (ev, pv, ac, progress)):
            continue
        features = {
            "planned_duration_days": planned_duration,
            "planned_budget_bac": bac,
            "elapsed_ratio": (status_date - planned_start).total_seconds() / (planned_duration * 86400),
            "percent_complete": progress,
            "ev": ev,
            "pv": pv,
            "ac": ac,
            "spi": ev / pv if pv > 0 else None,
            "cpi": ev / ac if ac > 0 else None,
            "sv": ev - pv,
            "cv": ev - ac,
        }
        if any(_number(features[name]) is None for name in SNAPSHOT_FEATURES):
            continue
        snapshots.append({
            "project_id": project_id,
            "snapshot_time": status_date.isoformat(),
            **features,
            "delay_days_raw": delay,
            "cost_overrun_ratio": overrun,
        })
    workbook.close()
    return snapshots


def prepare_empirical_data() -> dict[str, Any]:
    _extract_archive()
    outcomes = _read_project_outcomes()
    excel_dir = DATASET_DIR / "Excel"
    project_paths = {
        match.group(1).upper(): path
        for path in excel_dir.glob("*.xlsx")
        if (match := FILE_PROJECT_CODE.match(path.name))
    }
    project_rows = []
    snapshots = []
    for project_id, outcome in sorted(outcomes.items()):
        path = project_paths.get(project_id)
        if path is None or not _has_baseline_schedule(path):
            continue
        planned_duration = outcome["planned_duration_days"]
        bac = outcome["planned_budget_bac"]
        actual_duration = outcome["actual_duration_days"]
        actual_cost = outcome["actual_cost"]
        delay = actual_duration - planned_duration if planned_duration and planned_duration > 0 and actual_duration is not None and actual_duration >= 0 else None
        cost = (actual_cost - bac) / bac if bac is not None and bac > 0 and actual_cost is not None and actual_cost >= 0 else None
        project_rows.append({
            "project_id": project_id,
            "planned_duration_days": planned_duration,
            "planned_budget_bac": bac,
            "delay_days_raw": delay,
            "cost_overrun_ratio": cost,
        })
        start, finish = _baseline_dates(path)
        snapshots.extend(_safe_snapshots(path, project_id, outcome, start, finish))

    interim_dir = ROOT / "data" / "interim" / "dslib"
    processed_dir = ROOT / "data" / "processed" / "dslib"
    interim_dir.mkdir(parents=True, exist_ok=True)
    processed_dir.mkdir(parents=True, exist_ok=True)
    project_frame = pd.DataFrame(project_rows)
    snapshot_frame = pd.DataFrame(snapshots)
    project_path = interim_dir / "project_samples.csv"
    snapshot_path = interim_dir / "temporal_snapshots.csv"
    processed_path = processed_dir / "model_samples.csv"
    project_frame.to_csv(project_path, index=False)
    snapshot_frame.to_csv(snapshot_path, index=False)
    project_frame.to_csv(processed_path, index=False)

    source_manifest_path = RAW_DIR / "dataset_manifest.json"
    source_manifest = json.loads(source_manifest_path.read_text(encoding="utf-8")) if source_manifest_path.exists() else {}
    feature_columns = PROJECT_FEATURES
    complete_features = project_frame[PROJECT_FEATURES].notna().all(axis=1)
    prepared = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": "DSLIB / PSLIB release v3.4",
        "source_manifest_sha256": hashlib.sha256(source_manifest_path.read_bytes()).hexdigest() if source_manifest_path.exists() else None,
        "raw_archive_sha256": source_manifest.get("files", [{}])[0].get("sha256"),
        "project_count": len(project_rows),
        "project_samples": {
            "rows": len(project_rows),
            "delay_eligible_projects": int((complete_features & project_frame["delay_days_raw"].notna()).sum()),
            "cost_eligible_projects": int((complete_features & project_frame["cost_overrun_ratio"].notna()).sum()),
            "features": feature_columns,
            "targets": ["delay_days_raw", "cost_overrun_ratio"],
        },
        "temporal_snapshots": {
            "rows": len(snapshot_frame),
            "projects": int(snapshot_frame["project_id"].nunique()) if not snapshot_frame.empty else 0,
            "features": SNAPSHOT_FEATURES,
            "target_columns": ["delay_days_raw", "cost_overrun_ratio"],
            "cutoff": "tracking status date strictly before baseline project finish",
        },
        "files": [
            {"path": str(project_path.relative_to(ROOT)), "sha256": hashlib.sha256(project_path.read_bytes()).hexdigest()},
            {"path": str(snapshot_path.relative_to(ROOT)), "sha256": hashlib.sha256(snapshot_path.read_bytes()).hexdigest()},
            {"path": str(processed_path.relative_to(ROOT)), "sha256": hashlib.sha256(processed_path.read_bytes()).hexdigest()},
        ],
        "excluded_features": ["project name", "submitted by", "analysis-sheet final metrics", "actual duration", "actual cost"],
        "notes": [
            "Project-level rows use only baseline planned duration and BAC as features.",
            "Targets are final actual duration minus planned duration (days), and (final actual cost - BAC) / BAC.",
            "Negative delay and cost-overrun values are preserved.",
            "Temporal rows require observed project summary EV/PV/AC/progress and tracking time before planned finish.",
            "No synthetic observations, inferred categories, or Health labels are created.",
        ],
    }
    manifest_path = interim_dir / "preparation_manifest.json"
    manifest_path.write_text(json.dumps(prepared, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return prepared
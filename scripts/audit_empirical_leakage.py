from __future__ import annotations

import json
import math
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from audit_dslib import DATASET_DIR, REPORT_DIR, ROOT, extract_archive, read_project_outcomes


FEATURE_NAMES = [
    "planned_duration_days",
    "planned_budget_bac",
    "elapsed_ratio",
    "percent_complete",
    "ev",
    "pv",
    "ac",
    "spi",
    "cpi",
    "sv",
    "cv",
]
TARGET_NAMES = ["delay_days_raw", "cost_overrun_ratio"]


def numeric(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        return None
    return float(value)


def as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return None


def baseline_dates(path: Path) -> tuple[datetime | None, datetime | None]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Baseline Schedule" not in workbook.sheetnames:
        workbook.close()
        return None, None
    sheet = workbook["Baseline Schedule"]
    header_row = None
    start_index = end_index = id_index = None
    for row_number in range(1, min(sheet.max_row, 12) + 1):
        values = [sheet.cell(row=row_number, column=column).value for column in range(1, sheet.max_column + 1)]
        normalized = {str(value).strip().casefold(): index for index, value in enumerate(values) if value is not None}
        if "baseline start" in normalized and "baseline end" in normalized and "id" in normalized:
            header_row = row_number
            id_index = normalized["id"]
            start_index = normalized["baseline start"]
            end_index = normalized["baseline end"]
            break
    start_value = end_value = None
    if header_row is not None:
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if id_index is not None and len(row) > max(id_index, start_index or 0, end_index or 0) and row[id_index] == 0:
                start_value = as_datetime(row[start_index])
                end_value = as_datetime(row[end_index])
                break
    workbook.close()
    return start_value, end_value


def project_snapshots(path: Path, planned_start: datetime | None, planned_finish: datetime | None,
                       planned_duration: float | None, bac: float | None) -> tuple[list[dict[str, Any]], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    candidates: list[dict[str, Any]] = []
    post_finish_count = 0
    for sheet_name in workbook.sheetnames:
        if not sheet_name.casefold().startswith("project control"):
            continue
        sheet = workbook[sheet_name]
        status_date = as_datetime(sheet.cell(row=1, column=3).value)
        if status_date is None:
            continue
        headers = [sheet.cell(row=4, column=column).value for column in range(1, sheet.max_column + 1)]
        columns = {str(value).strip().casefold(): index for index, value in enumerate(headers) if value is not None}
        required = {
            "ev": columns.get("earned value (ev)"),
            "pv": columns.get("planned value (pv)"),
            "ac": columns.get("actual cost"),
            "progress": columns.get("percentage completed"),
        }
        if any(index is None for index in required.values()):
            continue
        summary = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if row and row[0] == 0:
                summary = row
                break
        if summary is None:
            continue
        ev = numeric(summary[required["ev"]])
        pv = numeric(summary[required["pv"]])
        ac = numeric(summary[required["ac"]])
        progress = numeric(summary[required["progress"]])
        if any(value is None for value in (ev, pv, ac, progress)):
            continue
        if planned_finish is None or status_date >= planned_finish:
            post_finish_count += 1
            continue
        if planned_start is None or planned_duration is None or planned_duration <= 0 or bac is None or bac <= 0:
            continue
        elapsed_ratio = (status_date - planned_start).total_seconds() / (planned_duration * 86400)
        row_features = {
            "planned_duration_days": planned_duration,
            "planned_budget_bac": bac,
            "elapsed_ratio": elapsed_ratio,
            "percent_complete": progress,
            "ev": ev,
            "pv": pv,
            "ac": ac,
            "spi": ev / pv if pv > 0 else None,
            "cpi": ev / ac if ac > 0 else None,
            "sv": ev - pv,
            "cv": ev - ac,
        }
        if all(numeric(row_features[name]) is not None for name in FEATURE_NAMES):
            candidates.append({"status_date": status_date.isoformat(), "features": row_features})
    workbook.close()
    return candidates, post_finish_count


def main() -> None:
    extract_archive()
    outcomes = read_project_outcomes(DATASET_DIR / "DSLIB_Analysis_Sheet.xlsx")
    excel_dir = DATASET_DIR / "Excel"
    project_paths = {
        path.name.split(" ", 1)[0].upper(): path
        for path in excel_dir.glob("*.xlsx")
    }
    candidates: dict[str, list[dict[str, Any]]] = {"delay": [], "cost": []}
    projects_with_rows: dict[str, set[str]] = {"delay": set(), "cost": set()}
    excluded_after_finish = 0
    checked_snapshots = 0
    for project_id, outcome in outcomes.items():
        path = project_paths.get(project_id)
        if path is None:
            continue
        planned_duration = numeric(outcome.get("planned_duration"))
        bac = numeric(outcome.get("bac"))
        actual_duration = numeric(outcome.get("actual_duration"))
        final_cost = numeric(outcome.get("final_cost"))
        delay_target = actual_duration - planned_duration if actual_duration is not None and planned_duration is not None else None
        cost_target = (final_cost - bac) / bac if final_cost is not None and bac is not None and bac > 0 else None
        planned_start, planned_finish = baseline_dates(path)
        snapshots, post_finish = project_snapshots(path, planned_start, planned_finish, planned_duration, bac)
        excluded_after_finish += post_finish
        checked_snapshots += len(snapshots) + post_finish
        for snapshot in snapshots:
            if delay_target is not None:
                candidates["delay"].append({"project_id": project_id, **snapshot, "target": delay_target})
                projects_with_rows["delay"].add(project_id)
            if cost_target is not None:
                candidates["cost"].append({"project_id": project_id, **snapshot, "target": cost_target})
                projects_with_rows["cost"].add(project_id)

    target_only = ["actual_duration", "final_cost", "delay_days_raw", "cost_overrun_ratio"]
    leakage_findings = [
        {
            "field": field,
            "role": "target_only",
            "used_as_feature": field in FEATURE_NAMES,
            "status": "clear" if field not in FEATURE_NAMES else "leakage_detected",
        }
        for field in target_only
    ]
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": "DSLIB / PSLIB release v3.4",
        "feature_names": FEATURE_NAMES,
        "target_names": TARGET_NAMES,
        "target_only_fields": target_only,
        "leakage_findings": leakage_findings,
        "post_planned_finish_snapshots_excluded": excluded_after_finish,
        "candidate_snapshots_checked": checked_snapshots,
        "snapshot_cutoff": "TP Status Date < baseline project finish; missing baseline dates are excluded.",
        "sample_counts": {
            "health": {"projects": 0, "snapshots": 0, "status": "not_available", "reason": "No consistent observed Health label; no threshold-derived label created."},
            "delay": {"projects": len(projects_with_rows["delay"]), "snapshots": len(candidates["delay"])},
            "cost": {"projects": len(projects_with_rows["cost"]), "snapshots": len(candidates["cost"])},
        },
        "status": "clear_after_filter" if all(item["status"] == "clear" for item in leakage_findings) else "blocked",
        "limitations": [
            "Only project-level TP summary rows with observed EV, PV, AC, and percent complete are eligible.",
            "Final duration and cost are targets only; they are not included in the feature vector.",
            "This audit does not establish that the remaining sample is sufficient for reliable generalization.",
        ],
    }
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    output = REPORT_DIR / "leakage_audit.json"
    output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    sample_path = REPORT_DIR / "temporal_sample_counts.json"
    sample_path.write_text(json.dumps(report["sample_counts"], indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "post_finish_excluded": excluded_after_finish, "sample_counts": report["sample_counts"], "report": str(output.relative_to(ROOT))}, indent=2))


if __name__ == "__main__":
    main()
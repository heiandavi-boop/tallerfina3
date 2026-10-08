#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath

import requests

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw"
MANIFEST_PATH = RAW / "dataset_manifest.json"
MENDELEY_ID = "2p5sz57wh2"
MENDELEY_VERSION = 2
MENDELEY_DOI = "10.17632/2p5sz57wh2.2"
MENDELEY_PAGE = f"https://data.mendeley.com/datasets/{MENDELEY_ID}/{MENDELEY_VERSION}"
ZENODO_API = "https://zenodo.org/api/records"
CHUNK_SIZE = 1024 * 1024


def _get_json(url: str, params: dict | None = None) -> dict:
    response = requests.get(url, params=params, timeout=60)
    response.raise_for_status()
    return response.json()


def _mendeley_downloads(metadata: dict, required_names: set[str]) -> dict[str, dict]:
    files = {item.get("filename"): item for item in metadata.get("files", [])}
    missing = required_names - files.keys()
    if missing:
        raise RuntimeError("La metadata Mendeley no incluye archivos: " + ", ".join(sorted(missing)))
    resolved = {}
    for filename in required_names:
        item = files[filename]
        download_url = (item.get("content_details") or {}).get("download_url")
        if not download_url:
            raise RuntimeError(f"Mendeley no publicó URL de descarga para {filename}")
        resolved[filename] = item
    return resolved


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(CHUNK_SIZE), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _md5(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(CHUNK_SIZE), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _valid_file(path: Path, filename: str, required_columns: tuple[str, ...] = ()) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError(f"Archivo ausente o vacío: {path}")
    suffix = filename.lower()
    if suffix.endswith(".csv"):
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            header = next(csv.reader(stream), None)
        _validate_columns(header, path, required_columns)
    elif suffix.endswith(".xlsx") or suffix.endswith(".zip"):
        if not zipfile.is_zipfile(path):
            raise ValueError(f"Formato ZIP/XLSX inválido: {path}")
        with zipfile.ZipFile(path) as archive:
            corrupt_member = archive.testzip()
            if corrupt_member:
                raise ValueError(f"Miembro ZIP corrupto {corrupt_member}: {path}")
            if suffix.endswith(".xlsx") and "xl/workbook.xml" not in archive.namelist():
                raise ValueError(f"No contiene un libro XLSX: {path}")
        if suffix.endswith(".xlsx") and required_columns:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook[workbook.sheetnames[0]]
                header = next(sheet.iter_rows(values_only=True), None)
                columns = [str(value or "") for value in header or ()]
                if len(columns) == 1 and "," in columns[0]:
                    columns = next(csv.reader([columns[0]]), [])
                _validate_columns(columns, path, required_columns)
            finally:
                workbook.close()
    elif suffix.endswith(".parquet"):
        with path.open("rb") as stream:
            if stream.read(4) != b"PAR1":
                raise ValueError(f"Firma Parquet inválida: {path}")
            stream.seek(-4, 2)
            if stream.read(4) != b"PAR1":
                raise ValueError(f"Footer Parquet inválido: {path}")
        import pyarrow.parquet as parquet

        schema = parquet.read_schema(path)
        if not schema.names or len(schema.names) != len(set(schema.names)):
            raise ValueError(f"Schema Parquet vacío o con nombres duplicados: {path}")


def _validate_columns(header, path: Path, required_columns: tuple[str, ...]) -> None:
    if not header or not any(str(column).strip() for column in header):
        raise ValueError(f"Tabla sin encabezado válido: {path}")
    normalized = {re.sub(r"[^a-z0-9]", "", str(column).lower()) for column in header}
    for group in required_columns:
        aliases = group.split("|")
        if not any(re.sub(r"[^a-z0-9]", "", alias) in normalized for alias in aliases):
            raise ValueError(f"Esquema inválido en {path}: falta uno de {aliases}")


def _cache_matches(previous: dict, dataset: str, filename: str, path: Path, fingerprint: str, expected_size: int | None, checksum: str | None) -> bool:
    old = next((row for row in previous.get("files", [])
                if row.get("dataset") == dataset and row.get("filename") == filename), None)
    if not old or not path.is_file() or path.stat().st_size == 0:
        return False
    if old.get("source_fingerprint") != fingerprint or old.get("validation_status") != "valid":
        return False
    if old.get("bytes") != path.stat().st_size or old.get("sha256") != _sha256(path):
        return False
    if expected_size is not None and path.stat().st_size != expected_size:
        return False
    if checksum:
        algorithm, _, expected_digest = checksum.partition(":")
        if algorithm == "sha256" and _sha256(path) != expected_digest:
            return False
        if algorithm == "md5" and _md5(path) != expected_digest:
            return False
    return True


def _download(url: str, destination: Path, expected_size: int | None, checksum: str | None) -> tuple[str, str]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".part")
    try:
        with requests.get(url, stream=True, timeout=(30, 180)) as response:
            response.raise_for_status()
            content_type = response.headers.get("Content-Type", "")
            with temporary.open("wb") as stream:
                for chunk in response.iter_content(chunk_size=CHUNK_SIZE):
                    if chunk:
                        stream.write(chunk)
            final_url = response.url
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise ValueError(f"Descarga vacía: {url}")
        _validate_content_type(destination.name, content_type)
        if expected_size is not None and temporary.stat().st_size != expected_size:
            raise ValueError(f"Tamaño incorrecto: esperado {expected_size}, recibido {temporary.stat().st_size}")
        if checksum:
            algorithm, _, expected_digest = checksum.partition(":")
            if algorithm == "sha256" and _sha256(temporary) != expected_digest:
                raise ValueError(f"Checksum SHA256 incorrecto: {destination.name}")
            if algorithm == "md5" and _md5(temporary) != expected_digest:
                raise ValueError(f"Checksum MD5 incorrecto: {destination.name}")
        temporary.replace(destination)
        return final_url, content_type
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _validate_content_type(filename: str, content_type: str) -> None:
    media_type = content_type.partition(";")[0].strip().lower()
    if not media_type or media_type in {"application/octet-stream", "binary/octet-stream"}:
        return
    if media_type in {"text/html", "application/json", "text/plain"}:
        if filename.lower().endswith((".xlsx", ".zip", ".parquet")) or media_type == "text/html":
            raise ValueError(f"Content-Type inesperado para {filename}: {content_type}")


def _record(dataset: str, landing_page: str, doi: str, record: str, filename: str,
            path: Path, url: str, content_type: str, license_name: str,
            synthetic: bool, fingerprint: str, status: str = "valid", **extra) -> dict:
    return {
        "dataset": dataset,
        "landing_page": landing_page,
        "doi": doi,
        "record/version": record,
        "resolved_download_url": url,
        "local_path": str(path.relative_to(ROOT)),
        "filename": filename,
        "bytes": path.stat().st_size,
        "sha256": _sha256(path),
        "downloaded_at": datetime.now(timezone.utc).isoformat(),
        "content_type": content_type,
        "license": license_name,
        "synthetic": synthetic,
        "validation_status": status,
        "source_fingerprint": fingerprint,
        **extra,
    }


def _download_one(previous: dict, rows: list[dict], *, dataset: str, landing_page: str,
                  doi: str, record: str, filename: str, path: Path, url: str,
                  license_name: str, synthetic: bool, fingerprint: str,
                  expected_size: int | None = None, checksum: str | None = None,
                  required_columns: tuple[str, ...] = ()) -> dict:
    if _cache_matches(previous, dataset, filename, path, fingerprint, expected_size, checksum):
        _valid_file(path, filename, required_columns)
        print(f"[OK] Using cached local dataset: {path.relative_to(ROOT)}")
        row = next(item for item in previous["files"]
                   if item.get("dataset") == dataset and item.get("filename") == filename)
        rows.append(row)
        return row
    print(f"[DOWNLOAD] {filename}")
    resolved_url, content_type = _download(url, path, expected_size, checksum)
    _valid_file(path, filename, required_columns)
    row = _record(dataset, landing_page, doi, record, filename, path, resolved_url,
                  content_type, license_name, synthetic, fingerprint,
                  expected_size=expected_size, checksum=checksum)
    rows.append(row)
    print(f"[OK] {path.relative_to(ROOT)} | {row['bytes']} bytes | SHA256 {row['sha256']}")
    return row


def download_mendeley(previous: dict, rows: list[dict]) -> None:
    print("\nMendeley")
    metadata_url = f"https://data.mendeley.com/public-api/datasets/{MENDELEY_ID}"
    metadata = _get_json(metadata_url, {"version": MENDELEY_VERSION, "fields": "*"})
    if int(metadata.get("version", 0)) != MENDELEY_VERSION or metadata.get("doi", {}).get("id") != MENDELEY_DOI:
        raise RuntimeError("La metadata de Mendeley no coincide con la versión/DOI canónicos.")
    print(f"[CHECK] Version {MENDELEY_VERSION}")
    filenames = {
        "project_risk_raw_dataset.csv": ("projectid|project_id|project",),
        "simulated_project_ev_metrics_final.csv.xlsx": ("projectid|project_id|project", "pv|plannedvalue", "ev|earnedvalue", "ac|actualcost"),
        "project_delay_cost_overrun_summary_final.csv.xlsx": ("projectid|project_id|project",),
    }
    resolved_files = _mendeley_downloads(metadata, set(filenames))
    license_name = (metadata.get("data_licence") or {}).get("full_name", "Unknown")
    for filename, required_columns in filenames.items():
        file_info = resolved_files[filename]
        details = file_info.get("content_details") or {}
        checksum = f"sha256:{details['sha256_hash']}" if details.get("sha256_hash") else None
        expected_size = details.get("size") or file_info.get("size")
        fingerprint = hashlib.sha256(
            f"{MENDELEY_DOI}|v{MENDELEY_VERSION}|{filename}|{expected_size}|{checksum}".encode()
        ).hexdigest()
        _download_one(previous, rows, dataset="mendeley_evm", landing_page=MENDELEY_PAGE,
                      doi=MENDELEY_DOI, record="version 2", filename=filename,
                      path=RAW / "mendeley" / filename, url=details["download_url"],
                      license_name=license_name, synthetic=True, fingerprint=fingerprint,
                      expected_size=expected_size, checksum=checksum,
                      required_columns=required_columns)


def _zenodo_files(record_id: int) -> tuple[dict, list[dict]]:
    metadata = _get_json(f"{ZENODO_API}/{record_id}")
    return metadata, metadata.get("files", [])


def _zenodo_file_url(file_metadata: dict) -> str:
    url = (file_metadata.get("links") or {}).get("self")
    if not url:
        raise RuntimeError(f"Zenodo no publicó enlace de descarga para {file_metadata.get('key')}")
    return url


def _zenodo_license(metadata: dict) -> str:
    license_data = metadata.get("metadata", {}).get("license") or {}
    if isinstance(license_data, dict):
        return license_data.get("id") or license_data.get("title") or "Unknown"
    return str(license_data)


def _download_zenodo_file(previous: dict, rows: list[dict], dataset: str, record_id: int,
                          landing_page: str, doi: str, file_info: dict,
                          license_name: str,
                          destination_name: str | None = None) -> Path:
    filename = destination_name or file_info["key"]
    path = RAW / dataset.removesuffix("_dataset") / filename
    record_meta = file_info.get("checksum", "")
    fingerprint = hashlib.sha256(
        f"{record_id}|{file_info.get('key')}|{file_info.get('size')}|{record_meta}".encode()
    ).hexdigest()
    _download_one(previous, rows, dataset=dataset, landing_page=landing_page, doi=doi,
                  record=f"Zenodo record {record_id}", filename=filename, path=path,
                  url=_zenodo_file_url(file_info), license_name=license_name,
                  synthetic=False, fingerprint=fingerprint,
                  expected_size=file_info.get("size"), checksum=record_meta or None)
    return path


def download_itemlet(previous: dict, rows: list[dict]) -> None:
    print("\nItemlet")
    record_id = 19411554
    landing_page = f"https://zenodo.org/records/{record_id}"
    doi = "10.5281/zenodo.19411554"
    metadata, available = _zenodo_files(record_id)
    by_name = {item.get("key"): item for item in available}
    preferred = "itemlet_dataset.parquet" if "itemlet_dataset.parquet" in by_name else "itemlet_dataset.csv"
    required = [preferred, "data_dictionary.csv", "project_domain.csv", "project_summary.csv"]
    missing = [name for name in required if name not in by_name]
    if missing:
        raise RuntimeError("Archivos Itemlet requeridos ausentes en metadata: " + ", ".join(missing))
    license_name = _zenodo_license(metadata)
    for filename in required:
        item = by_name[filename]
        _download_zenodo_file(previous, rows, "itemlet", record_id, landing_page, doi, item, license_name)


def _safe_extract(archive_path: Path, target: Path) -> list[Path]:
    extracted: list[Path] = []
    with zipfile.ZipFile(archive_path) as archive:
        for member in archive.infolist():
            relative = PurePosixPath(member.filename)
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError(f"Ruta insegura en ZIP: {member.filename}")
            destination = target.joinpath(*relative.parts)
            if member.is_dir():
                destination.mkdir(parents=True, exist_ok=True)
                continue
            destination.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(member) as source, destination.open("wb") as output:
                for chunk in iter(lambda: source.read(CHUNK_SIZE), b""):
                    output.write(chunk)
            extracted.append(destination)
    return extracted


def _collaboration_feature_signals(member_name: str) -> list[str]:
    normalized = member_name.lower()
    if not normalized.endswith(".csv"):
        return []
    if "/collaboration_metrics/" in normalized:
        if "teamm" in Path(member_name).name.lower():
            return ["team_size", "team_expertise", "core_contribution"]
        if "communicationmetrics" in Path(member_name).name.lower():
            return ["issue_interactions", "pull_request_interactions"]
    if "/delivery_indicators/" in normalized:
        filename = Path(member_name).name.lower()
        if filename.startswith("deliverySize_".lower()):
            return ["delivery_size"]
        if filename.startswith("deliveryFrequency_".lower()):
            return ["delivery_frequency"]
        if filename.startswith("changeLeadTime_".lower()):
            return ["change_lead_time"]
    return []


def download_collaboration(previous: dict, rows: list[dict]) -> None:
    print("\nCollaboration")
    record_id = 15681547
    landing_page = f"https://zenodo.org/records/{record_id}"
    doi = "10.5281/zenodo.15681547"
    metadata, available = _zenodo_files(record_id)
    archive = next((item for item in available if item.get("key") == "dataset.zip"), None)
    if archive is None:
        raise RuntimeError("Zenodo record Collaboration no publica dataset.zip.")
    archive_path = _download_zenodo_file(previous, rows, "collaboration", record_id,
                                         landing_page, doi, archive, _zenodo_license(metadata))
    extracted_paths = _safe_extract(archive_path, RAW / "collaboration")
    archive_row = next(item for item in rows if item["dataset"] == "collaboration" and item["filename"] == "dataset.zip")
    archive_row["extracted_files"] = []
    archive_row["files_used"] = []
    archive_row["eligible_feature_files"] = []
    archive_row["feature_usage_note"] = (
        "No se unen con Mendeley ni se usan como labels/features supervisadas sin una "
        "preparación con equivalencia temporal y unidad validada."
    )
    for extracted in extracted_paths:
        _valid_file(extracted, extracted.name)
        relative = extracted.relative_to(ROOT)
        signals = _collaboration_feature_signals(extracted.as_posix())
        archive_row["extracted_files"].append(str(relative))
        if signals:
            archive_row["eligible_feature_files"].append(str(relative))
        rows.append({
            "dataset": "collaboration", "landing_page": landing_page, "doi": doi,
            "record/version": f"Zenodo record {record_id}", "resolved_download_url": archive_row["resolved_download_url"],
            "local_path": str(relative), "filename": extracted.name,
            "bytes": extracted.stat().st_size, "sha256": _sha256(extracted),
            "downloaded_at": archive_row["downloaded_at"], "content_type": "application/zip member",
            "license": _zenodo_license(metadata),
            "synthetic": False, "validation_status": "valid", "source_fingerprint": archive_row["source_fingerprint"],
            "source_archive": "data/raw/collaboration/dataset.zip", "used_for_features": False,
            "eligible_feature_signals": signals,
            "selection_status": "eligible_not_joined" if signals else "not_selected",
        })
    print(
        f"[OK] dataset.zip validado y extraído | {len(extracted_paths)} archivos | "
        f"{len(archive_row['eligible_feature_files'])} CSV de señales elegibles, sin JOIN"
    )


SQUAD_TABLES = ("projects", "commits", "issues", "releases", "github_metrics", "process_metrics")


def download_squad(previous: dict, rows: list[dict], max_bytes: int) -> None:
    print("\nSQuaD")
    record_id = 17566691
    landing_page = f"https://zenodo.org/records/{record_id}"
    doi = "10.5281/zenodo.17566691"
    metadata, available = _zenodo_files(record_id)
    print("[INFO] Resolviendo subset requerido...")
    selected = []
    for item in available:
        key = item.get("key", "").lower()
        if key.endswith((".csv", ".parquet")) and any(table in key for table in SQUAD_TABLES):
            if int(item.get("size", 0)) <= max_bytes:
                selected.append(item)
    if not selected:
        print("[INFO] El record no expone tablas descargables aptas; SQuaD queda omitido (dataset opcional).")
        return
    for item in selected:
        _download_zenodo_file(previous, rows, "squad", record_id, landing_page, doi,
                              item, _zenodo_license(metadata))
    print(f"[OK] {len(selected)} archivo(s) SQuaD seleccionados por nombre/tamaño")


def _write_manifest(rows: list[dict], source_metadata: dict) -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    manifest = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_metadata": source_metadata,
        "files": rows,
    }
    temporary = MANIFEST_PATH.with_name(MANIFEST_PATH.name + ".part")
    temporary.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(MANIFEST_PATH)


def validate_local_datasets() -> dict:
    if not MANIFEST_PATH.is_file():
        raise RuntimeError(f"Falta manifest local: {MANIFEST_PATH}")
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    rows = manifest.get("files", [])
    by_path = {row.get("local_path"): row for row in rows}
    required = [
        "data/raw/mendeley/project_risk_raw_dataset.csv",
        "data/raw/mendeley/simulated_project_ev_metrics_final.csv.xlsx",
        "data/raw/mendeley/project_delay_cost_overrun_summary_final.csv.xlsx",
    ]
    missing = [path for path in required if path not in by_path]
    if missing:
        raise RuntimeError("Mendeley no contiene los archivos obligatorios: " + ", ".join(missing))
    required_rows = [by_path[path] for path in required]
    if any(row.get("dataset") != "mendeley_evm" or row.get("synthetic") is not True
           or row.get("doi") != MENDELEY_DOI or row.get("record/version") != "version 2"
           for row in required_rows):
        raise RuntimeError("Los archivos principales no acreditan Mendeley v2 sintético.")
    mendeley_columns = {
        "data/raw/mendeley/project_risk_raw_dataset.csv": ("projectid|project_id|project",),
        "data/raw/mendeley/simulated_project_ev_metrics_final.csv.xlsx": (
            "projectid|project_id|project", "pv|plannedvalue", "ev|earnedvalue", "ac|actualcost"
        ),
        "data/raw/mendeley/project_delay_cost_overrun_summary_final.csv.xlsx": ("projectid|project_id|project",),
    }
    optional_errors = []
    for row in rows:
        relative = row.get("local_path", "")
        try:
            path = (ROOT / relative).resolve()
            if ROOT.resolve() not in path.parents:
                raise RuntimeError(f"Ruta fuera del workspace en manifest: {relative}")
            if row.get("validation_status") != "valid" or not path.is_file():
                raise RuntimeError(f"Archivo local ausente/no validado: {relative}")
            if path.stat().st_size != row.get("bytes") or _sha256(path) != row.get("sha256"):
                raise RuntimeError(f"Tamaño/SHA256 no coincide con manifest: {relative}")
            _valid_file(path, row.get("filename", path.name), mendeley_columns.get(relative, ()))
        except (OSError, ValueError, RuntimeError) as error:
            if row.get("dataset") == "mendeley_evm":
                raise
            optional_errors.append(str(error))
    if optional_errors:
        manifest.setdefault("warnings", []).extend(optional_errors)
    return manifest


def _optional_download(name: str, action, rows: list[dict], statuses: dict, warnings: list[str]) -> None:
    first_row = len(rows)
    try:
        action()
        relevant_rows = rows[first_row:]
        statuses[name] = "available" if relevant_rows else "unavailable"
        if not relevant_rows:
            raise RuntimeError("metadata no resolvió archivos descargables")
    except Exception as error:
        del rows[first_row:]
        statuses[name] = "unavailable"
        warning = f"{name}: {error}"
        warnings.append(warning)
        print(f"[WARN] Optional dataset unavailable: {name} ({error})")


def main() -> int:
    parser = argparse.ArgumentParser(description="Descarga y valida datasets PRUNIN en data/raw/.")
    parser.add_argument("--include-squad", action="store_true", help="Descarga solo tablas SQuaD relevantes y bajo el límite configurado.")
    parser.add_argument("--squad-max-bytes", type=int, default=1_000_000_000)
    parser.add_argument("--validate-only", action="store_true", help="Valida únicamente el manifest y archivos locales, sin acceso a red.")
    args = parser.parse_args()

    print("# PRUNIN AI Core - Dataset Downloader")
    if args.validate_only:
        manifest = validate_local_datasets()
        print(f"[OK] Archivos locales y hashes válidos: {len(manifest['files'])}")
        return 0
    print("[3] Resolviendo metadata oficial de datasets...")
    previous = json.loads(MANIFEST_PATH.read_text(encoding="utf-8")) if MANIFEST_PATH.is_file() else {}
    rows: list[dict] = []
    optional_status = {}
    warnings = []
    print("[4] Descargando datasets a data/raw/...")
    download_mendeley(previous, rows)
    optional_status["itemlet"] = "available"
    optional_status["collaboration"] = "available"
    _optional_download("itemlet", lambda: download_itemlet(previous, rows), rows, optional_status, warnings)
    _optional_download("collaboration", lambda: download_collaboration(previous, rows), rows, optional_status, warnings)
    if args.include_squad:
        _optional_download("squad", lambda: download_squad(previous, rows, args.squad_max_bytes),
                           rows, optional_status, warnings)
    else:
        optional_status["squad"] = "not_requested"
        warnings.append("SQuaD not requested; large source intentionally omitted.")
        print("\nSQuaD\n[WARN] Optional dataset unavailable: not requested (use --include-squad for a bounded subset).")
    source_metadata = {
        "mendeley_version": MENDELEY_VERSION,
        "zenodo_records": {"itemlet": 19411554, "collaboration": 15681547,
                           "squad": 17566691 if args.include_squad else None},
        "optional_datasets": optional_status,
    }
    _write_manifest(rows, source_metadata)
    local_manifest = validate_local_datasets()
    local_manifest["warnings"] = warnings + local_manifest.get("warnings", [])
    _write_manifest(rows, {**source_metadata, "warnings": local_manifest["warnings"]})
    print(f"\n[OK] Manifest: {MANIFEST_PATH.relative_to(ROOT)} | {len(rows)} entradas")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (requests.RequestException, OSError, ValueError, RuntimeError, KeyError) as error:
        print(f"[ERROR] {error}", file=sys.stderr)
        raise SystemExit(1)
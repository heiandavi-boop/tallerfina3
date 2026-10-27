from __future__ import annotations

import json
import math
import re
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw" / "dslib"
ARCHIVE = RAW_DIR / "DSLIB3.4.zip"
EXTRACTED_DIR = RAW_DIR / "extracted"
DATASET_DIR = EXTRACTED_DIR / "DSLIB 3.4"
REPORT_DIR = ROOT / "reports" / "empirical"
PROJECT_CODE = re.compile(r"^C\d{4}-\d{2}$", re.IGNORECASE)
FILE_PROJECT_CODE = re.compile(r"^(C\d{4}-\d{2})", re.IGNORECASE)


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def has_value(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip().casefold() not in {"", "n/a", "na", "-", "none"}
    return True


def extract_archive() -> None:
    if not ARCHIVE.exists():
        raise FileNotFoundError(f"No existe {ARCHIVE.relative_to(ROOT)}. Ejecute scripts/download_dslib.py.")
    if DATASET_DIR.exists():
        return
    EXTRACTED_DIR.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(ARCHIVE) as archive:
        for member in archive.infolist():
            destination = (EXTRACTED_DIR / member.filename).resolve()
            if not destination.is_relative_to(EXTRACTED_DIR.resolve()):
                raise ValueError(f"Ruta inválida dentro del ZIP: {member.filename}")
        archive.extractall(EXTRACTED_DIR)


def read_project_outcomes(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["DSLIB"]
    projects: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        raw_code = row[0] if row else None
        if not isinstance(raw_code, str) or not PROJECT_CODE.fullmatch(raw_code.strip()):
            continue
        code = raw_code.strip().upper()
        projects[code] = {
            "planned_duration": row[11] if len(row) > 11 else None,
            "bac": row[12] if len(row) > 12 else None,
            "actual_duration": row[82] if len(row) > 82 else None,
            "final_cost": row[83] if len(row) > 83 else None,
        }
    workbook.close()
    return projects


def read_project_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    result: dict[str, Any] = {
        "excel_file": True,
        "baseline": "Baseline Schedule" in workbook.sheetnames,
        "actuals": False,
        "evm": False,
        "pv": False,
        "ev": False,
        "ac": False,
        "temporal": False,
        "evm_snapshots": 0,
    }
    control_sheets = [name for name in workbook.sheetnames if name.casefold().startswith("project control")]
    for sheet_name in control_sheets:
        sheet = workbook[sheet_name]
        status_value = sheet.cell(row=1, column=3).value
        if isinstance(status_value, (datetime, date)):
            result["temporal"] = True

        headers = [sheet.cell(row=4, column=column).value for column in range(1, sheet.max_column + 1)]
        columns = {str(value).strip().casefold(): index for index, value in enumerate(headers) if value is not None}
        ev_index = columns.get("earned value (ev)")
        pv_index = columns.get("planned value (pv)")
        ac_index = columns.get("actual cost")
        duration_index = columns.get("actual duration")
        actual_start_index = columns.get("actual start")
        snapshot_values: dict[str, Any] = {"ev": None, "pv": None, "ac": None}

        for row in sheet.iter_rows(min_row=5, values_only=True):
            if actual_start_index is not None and actual_start_index < len(row) and has_value(row[actual_start_index]):
                result["actuals"] = True
            if duration_index is not None and duration_index < len(row) and has_value(row[duration_index]):
                result["actuals"] = True
            if ac_index is not None and ac_index < len(row) and has_value(row[ac_index]):
                result["actuals"] = True
            if row and row[0] == 0:
                for key, index in (("ev", ev_index), ("pv", pv_index), ("ac", ac_index)):
                    if index is not None and index < len(row) and is_number(row[index]):
                        snapshot_values[key] = row[index]
            for key, index in (("ev", ev_index), ("pv", pv_index), ("ac", ac_index)):
                if index is not None and index < len(row) and is_number(row[index]):
                    result[key] = True

        if all(is_number(snapshot_values[key]) for key in ("ev", "pv", "ac")):
            result["evm_snapshots"] += 1

    result["evm"] = result["ev"] and result["pv"] and result["ac"]
    workbook.close()
    return result


def count_project_files() -> dict[str, int]:
    counts: dict[str, int] = {}
    for path in DATASET_DIR.rglob("*"):
        if not path.is_file() or "__MACOSX" in path.parts or path.name == ".DS_Store":
            continue
        match = FILE_PROJECT_CODE.match(path.name)
        if match:
            code = match.group(1).upper()
            counts[code] = counts.get(code, 0) + 1
    return counts


def metric_row(name: str, count: int, total: int) -> dict[str, Any]:
    return {"variable": name, "projects_available": count, "coverage_percent": round(count * 100 / total, 2) if total else 0.0}


def main() -> None:
    extract_archive()
    analysis_path = DATASET_DIR / "DSLIB_Analysis_Sheet.xlsx"
    excel_dir = DATASET_DIR / "Excel"
    if not analysis_path.exists() or not excel_dir.exists():
        raise FileNotFoundError("El ZIP no tiene la estructura esperada (analysis sheet y Excel/).")

    outcomes = read_project_outcomes(analysis_path)
    project_paths = {
        match.group(1).upper(): path
        for path in excel_dir.glob("*.xlsx")
        if (match := FILE_PROJECT_CODE.match(path.name))
    }
    files_per_project = count_project_files()
    project_ids = sorted(set(outcomes) | set(project_paths))
    records: dict[str, dict[str, Any]] = {}
    for project_id in project_ids:
        outcome = outcomes.get(project_id, {})
        workbook_info = read_project_workbook(project_paths[project_id]) if project_id in project_paths else {
            "excel_file": False, "baseline": False, "actuals": False, "evm": False,
            "pv": False, "ev": False, "ac": False, "temporal": False, "evm_snapshots": 0,
        }
        planned_duration = outcome.get("planned_duration")
        bac = outcome.get("bac")
        actual_duration = outcome.get("actual_duration")
        final_cost = outcome.get("final_cost")
        record = {
            **workbook_info,
            "project_files": files_per_project.get(project_id, 0),
            "planned_duration": is_number(planned_duration) and planned_duration > 0,
            "bac": is_number(bac) and bac > 0,
            "actual_duration": is_number(actual_duration) and actual_duration >= 0,
            "final_cost": is_number(final_cost) and final_cost >= 0,
        }
        record["incomplete"] = not (
            record["baseline"]
            and record["actuals"]
            and record["planned_duration"]
            and record["actual_duration"]
            and record["bac"]
            and record["final_cost"]
        )
        record["delay_eligible"] = record["baseline"] and record["planned_duration"] and record["bac"] and record["actual_duration"]
        record["cost_eligible"] = record["baseline"] and record["planned_duration"] and record["bac"] and record["final_cost"]
        records[project_id] = record

    total = len(records)
    fields = [
        "project Excel workbook", "baseline schedule", "actual/control data", "complete EVM (EV/PV/AC)",
        "BAC > 0", "PV numeric", "EV numeric", "AC numeric", "planned duration > 0",
        "actual duration observed", "dated tracking information", "final cost observed", "incomplete project",
    ]
    record_keys = [
        "excel_file", "baseline", "actuals", "evm", "bac", "pv", "ev", "ac",
        "planned_duration", "actual_duration", "temporal", "final_cost", "incomplete",
    ]
    coverage = [metric_row(label, sum(bool(records[project][key]) for project in records), total) for label, key in zip(fields, record_keys)]

    delay_projects = [project for project, record in records.items() if record["delay_eligible"]]
    cost_projects = [project for project, record in records.items() if record["cost_eligible"]]
    counts = {
        "health": {"projects": 0, "snapshots": 0, "status": "not_available", "reason": "No consistent observed Health/final-status label is defined by the publisher."},
        "delay": {"projects": len(delay_projects), "snapshots": 0, "target": "analysis-sheet actual Duration minus planned PD (days); candidate only"},
        "cost": {"projects": len(cost_projects), "snapshots": 0, "target": "(analysis-sheet actual Cost - BAC) / BAC; candidate only"},
    }
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": "DSLIB / PSLIB release v3.4",
        "source_url": "https://github.com/MarioVanhoucke/DSLIB-Dynamic-Scheduling-Empirical-Project-Library/releases/tag/DSLIB_v3.4",
        "license": "Not specified in publisher repository/release metadata",
        "project_count": total,
        "analysis_sheet_project_count": len(outcomes),
        "project_workbook_count": len(project_paths),
        "project_file_counts": {
            "projects_with_at_least_one_file": sum(count > 0 for count in files_per_project.values()),
            "minimum_files_per_project": min(files_per_project.values(), default=0),
            "maximum_files_per_project": max(files_per_project.values(), default=0),
        },
        "coverage": coverage,
        "sample_counts": counts,
        "temporal_evm_snapshot_candidates_before_leakage_filter": {
            "delay": sum(records[project]["evm_snapshots"] for project in delay_projects),
            "cost": sum(records[project]["evm_snapshots"] for project in cost_projects),
        },
        "inclusion_criteria": {
            "health": "Blocked: no publisher-defined, consistent Health label; do not infer one from arbitrary thresholds.",
            "delay": "Baseline workbook plus positive planned duration and observed nonnegative final actual duration; raw delay may be negative.",
            "cost": "Baseline workbook plus BAC > 0 and observed nonnegative final cost; preserve negative overrun as savings.",
            "snapshots": "Counts here are descriptive only. No snapshot is training-eligible until leakage audit verifies status date < planned finish and observed EVM values.",
        },
        "limitations": [
            "The analysis sheet includes final outcomes and summary metrics; these must never be used as snapshot features.",
            "Project-control sheets vary in availability and completeness.",
            "No Health label is inferred and no model is trained by this audit.",
            "The publisher's repository does not state a data license.",
        ],
    }

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    json_path = REPORT_DIR / "dslib_audit.json"
    json_path.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    lines = [
        "# DSLIB v3.4 Dataset Audit",
        "",
        f"- Projects found across analysis/workbooks: **{total}**",
        f"- Analysis-sheet projects: **{len(outcomes)}**",
        f"- Per-project Excel workbooks: **{len(project_paths)}**",
        "- License: **not specified by publisher**",
        "- Health training: **blocked** until a defensible, pre-registered label is available",
        "",
        "## Coverage",
        "",
        "| Variable | Projects available | Coverage |",
        "|---|---:|---:|",
    ]
    lines.extend(f"| {item['variable']} | {item['projects_available']} | {item['coverage_percent']:.2f}% |" for item in coverage)
    lines.extend([
        "",
        "## Candidate Sample Counts",
        "",
        "| Target | Baseline projects | Snapshots used for training |",
        "|---|---:|---:|",
    ])
    lines.extend(f"| {target.title()} | {values['projects']} | {values['snapshots']} |" for target, values in counts.items())
    lines.extend([
        "",
        "Delay candidate: actual duration minus planned duration, in days; negative values mean early completion.",
        "Cost candidate: (actual cost - BAC) / BAC; negative values mean below-budget execution.",
        "The temporal EVM candidate count is reported separately; the initial empirical models use one baseline row per project.",
        "",
        "## Limitations",
        "",
    ])
    lines.extend(f"- {item}" for item in report["limitations"])
    md_path = REPORT_DIR / "dslib_audit.md"
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    counts_path = REPORT_DIR / "model_sample_counts.json"
    counts_path.write_text(json.dumps(counts, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"project_count": total, "coverage": coverage, "sample_counts": counts, "report": str(json_path.relative_to(ROOT)), "markdown": str(md_path.relative_to(ROOT))}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()